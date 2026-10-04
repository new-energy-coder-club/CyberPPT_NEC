#!/usr/bin/env python3
"""CyberPPT Stage 3: generate editable NEC proposal PPTX with real photos."""

from __future__ import annotations

import json
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
RgbColor = RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------
OUT_DIR = Path("D:/Project_env/CyberPPT_NEC/output")
ASSET_DIR = OUT_DIR / "assets"
PPTX_PATH = OUT_DIR / "NEC_Corporate_Partnership_Proposal.pptx"
MANIFEST_PATH = OUT_DIR / "slide_manifest.json"
QA_PATH = OUT_DIR / "visual_qa_gate.json"

SLIDE_WIDTH_IN = 13.333
SLIDE_HEIGHT_IN = 7.5

# Ivory + Deep Blue palette
COLORS = {
    "ivory": RgbColor(0xF8, 0xF6, 0xEF),
    "ivory_dark": RgbColor(0xEF, 0xED, 0xE4),
    "deep_blue": RgbColor(0x0A, 0x1F, 0x44),
    "accent_blue": RgbColor(0x1E, 0x3A, 0x5F),
    "light_blue": RgbColor(0x3C, 0x6E, 0xA8),
    "highlight": RgbColor(0xC9, 0xA2, 0x27),  # muted gold
    "white": RgbColor(0xFF, 0xFF, 0xFF),
    "black": RgbColor(0x1A, 0x1A, 0x1A),
    "gray": RgbColor(0x66, 0x66, 0x66),
}

# Source photos (from user provided directory)
PHOTO_SRC = Path("D:/OneDrive/xwechat_files/wxid_bmhsmkv5h72v12_e0e7/temp/RWTemp/2026-06/9e20f478899dc29eb19741386f9343c8")
PHOTOS = {
    "campus_gate": "4a5d59c6d0c82f53bf741fd65e37428a.jpg",
    "team_group": "55c3d4d71123ec65470b342cc74395b4.jpg",
    "lab_work": "ed23381904ec4e0072c1c35d9b5211e1.jpg",
    "robot_team": "dc3f6c4d7be0db88c1dbfe5cdfae97e9.jpg",
    "night_debug": "558efea6363b0b93be49c347144f2806.jpg",
    "trae_award": "413bb85db818ddb57bdad9ce7ef5f2bb.jpg",
    "opc_event": "76969b7dd1c01b72d6c81ea1208da05a.jpg",
}

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def ensure_assets():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    copied = {}
    for key, filename in PHOTOS.items():
        src = PHOTO_SRC / filename
        dst = ASSET_DIR / filename
        if not dst.exists():
            shutil.copy(src, dst)
        copied[key] = dst
    return copied


def add_textbox(slide, left, top, width, height, text, *, font_size, bold=False,
                color=None, align=PP_ALIGN.LEFT, font_name="Microsoft YaHei",
                valign=MSO_ANCHOR.TOP, italic=False, line_spacing=1.2):
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.italic = italic
    p.font.name = font_name
    p.font.color.rgb = color or COLORS["black"]
    p.alignment = align
    p.line_spacing = line_spacing
    return shape


def add_rounded_rect(slide, left, top, width, height, fill_color, *, line_color=None, line_width=0):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()
    return shape


def add_rect(slide, left, top, width, height, fill_color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    return shape


def add_line(slide, x1, y1, x2, y2, color, width_pt=1):
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    connector.line.color.rgb = color
    connector.line.width = Pt(width_pt)
    return connector


def add_footer(slide, page_number_text, footer_text="常州工学院 NEC 新能源开发者社区 · 企业合作提案"):
    # Left footer
    add_textbox(slide, 0.45, 7.05, 6.0, 0.3, footer_text, font_size=7.0, color=COLORS["gray"], valign=MSO_ANCHOR.MIDDLE)
    # Page number badge
    badge = add_rounded_rect(slide, 12.15, 6.95, 0.7, 0.35, COLORS["deep_blue"])
    add_textbox(slide, 12.15, 6.95, 0.7, 0.35, page_number_text, font_size=14.0, bold=True,
                color=COLORS["white"], align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def add_page_header(slide, badge_num, title, subtitle=""):
    # Badge
    badge = add_rounded_rect(slide, 0.45, 0.35, 0.55, 0.55, COLORS["deep_blue"])
    add_textbox(slide, 0.45, 0.35, 0.55, 0.55, badge_num, font_size=16.0, bold=True,
                color=COLORS["white"], align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    # Title
    add_textbox(slide, 1.15, 0.36, 11.5, 0.45, title, font_size=24.0, bold=True,
                color=COLORS["deep_blue"], valign=MSO_ANCHOR.MIDDLE)
    if subtitle:
        add_textbox(slide, 1.15, 0.82, 11.5, 0.28, subtitle, font_size=11.0,
                    color=COLORS["gray"], valign=MSO_ANCHOR.MIDDLE)


def add_so_what(slide, text):
    bar = add_rect(slide, 0.45, 6.55, 12.45, 0.22, COLORS["deep_blue"])
    add_textbox(slide, 0.55, 6.55, 1.2, 0.22, "SO WHAT", font_size=10.0, bold=True,
                color=COLORS["white"], valign=MSO_ANCHOR.MIDDLE)
    add_textbox(slide, 1.85, 6.55, 10.85, 0.22, text, font_size=10.0,
                color=COLORS["white"], valign=MSO_ANCHOR.MIDDLE)


def add_bullet_list(slide, left, top, width, height, items, font_size=10.5):
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = shape.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(font_size)
        p.font.name = "Microsoft YaHei"
        p.font.color.rgb = COLORS["black"]
        p.space_after = Pt(6)
    return shape


def add_kpi_card(slide, left, top, width, height, number, label):
    shape = add_rounded_rect(slide, left, top, width, height, COLORS["white"], line_color=COLORS["ivory_dark"], line_width=1)
    add_textbox(slide, left + 0.08, top + 0.08, width - 0.16, height * 0.45, number,
                font_size=22.0, bold=True, color=COLORS["deep_blue"], align=PP_ALIGN.CENTER,
                valign=MSO_ANCHOR.MIDDLE)
    add_textbox(slide, left + 0.08, top + height * 0.5, width - 0.16, height * 0.4, label,
                font_size=9.5, color=COLORS["gray"], align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.TOP)
    return shape


def add_table(slide, left, top, rows, cols, col_widths_in, row_height_in, data):
    table = slide.shapes.add_table(rows, cols, Inches(left), Inches(top), Inches(sum(col_widths_in)), Inches(rows * row_height_in)).table
    for i, w in enumerate(col_widths_in):
        table.columns[i].width = Inches(w)
    for r in range(rows):
        table.rows[r].height = Inches(row_height_in)
        for c in range(cols):
            cell = table.cell(r, c)
            cell.text = data[r][c]
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(11.0 if r == 0 else 9.5)
            p.font.name = "Microsoft YaHei"
            p.font.bold = (r == 0)
            p.font.color.rgb = COLORS["white"] if r == 0 else COLORS["black"]
            p.alignment = PP_ALIGN.CENTER if c > 0 else PP_ALIGN.LEFT
            if r == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = COLORS["deep_blue"]
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = COLORS["ivory"] if r % 2 == 0 else COLORS["white"]
    return table


# ---------------------------------------------------------------------------
# Slide builders
# ---------------------------------------------------------------------------
def build_slide_01(prs, photos):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    # Background ivory
    bg = add_rect(slide, 0, 0, SLIDE_WIDTH_IN, SLIDE_HEIGHT_IN, COLORS["ivory"])
    # Photo at left ~38% area
    slide.shapes.add_picture(str(photos["campus_gate"]), Inches(0), Inches(0), width=Inches(5.1), height=Inches(7.5))
    # Deep blue overlay panel
    add_rect(slide, 5.1, 0, 8.233, 7.5, COLORS["deep_blue"])
    # Decorative line
    add_line(slide, 6.0, 2.2, 11.8, 2.2, COLORS["highlight"], 2)
    # Title
    add_textbox(slide, 6.0, 2.45, 6.6, 1.2, "常州工学院 NEC", font_size=38.0, bold=True,
                color=COLORS["white"], valign=MSO_ANCHOR.TOP)
    add_textbox(slide, 6.0, 3.7, 6.6, 0.8, "新能源开发者社区", font_size=32.0, bold=True,
                color=COLORS["white"], valign=MSO_ANCHOR.TOP)
    add_textbox(slide, 6.0, 4.6, 6.6, 0.5, "企业合作提案 · 2026", font_size=14.0,
                color=COLORS["highlight"], valign=MSO_ANCHOR.TOP)
    add_textbox(slide, 6.0, 5.3, 6.6, 0.9,
                "面向政府园区与投资机构\n寻求资金支持与场地服务，共建高校新能源创新生态",
                font_size=11.5, color=COLORS["white"], valign=MSO_ANCHOR.TOP, line_spacing=1.3)
    add_textbox(slide, 6.0, 7.05, 6.6, 0.3, "常州工学院 NEC 新能源开发者社区", font_size=7.0,
                color=COLORS["white"], valign=MSO_ANCHOR.MIDDLE)


def build_slide_02(prs, photos):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, SLIDE_WIDTH_IN, SLIDE_HEIGHT_IN, COLORS["ivory"])
    add_page_header(slide, "01", "关于 NEC", "以学生为主体的新能源开发者社区，链接课堂、竞赛与产业")
    # Left text
    add_textbox(slide, 0.55, 1.45, 6.2, 0.6, "社区定位", font_size=13.0, bold=True,
                color=COLORS["deep_blue"], valign=MSO_ANCHOR.MIDDLE)
    add_bullet_list(slide, 0.55, 2.05, 6.0, 2.4, [
        "扎根常州工学院，聚焦新能源与智能机器人交叉领域",
        "覆盖机械设计、嵌入式电控、计算机视觉、新能源应用与开源运营",
        "通过赛事驱动真实项目，培养可交付的工程能力",
        "面向政府园区、投资机构与企业，提供产学研合作入口"
    ], font_size=10.5)
    # Right KPIs
    add_kpi_card(slide, 7.0, 1.45, 2.6, 1.2, "60+", "在册成员")
    add_kpi_card(slide, 9.9, 1.45, 2.6, 1.2, "5", "技术方向")
    add_kpi_card(slide, 7.0, 2.9, 2.6, 1.2, "国家级", "竞赛奖项")
    add_kpi_card(slide, 9.9, 2.9, 2.6, 1.2, "持续", "迭代周期")
    add_so_what(slide, "NEC 是一支能持续产出工程成果的学生战队，具备企业合作与产业落地的底层能力")
    add_footer(slide, "02")


def build_slide_03(prs, photos):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, SLIDE_WIDTH_IN, SLIDE_HEIGHT_IN, COLORS["ivory"])
    add_page_header(slide, "02", "团队与组织", "跨学科协作，形成完整研发闭环")
    # Group photo right
    slide.shapes.add_picture(str(photos["team_group"]), Inches(7.2), Inches(1.35), width=Inches(5.7), height=Inches(3.2))
    # Org blocks left
    orgs = [
        ("机械组", "结构 / 传动 / 加工", COLORS["deep_blue"]),
        ("电控组", "嵌入式 / 电机 / 通信", COLORS["accent_blue"]),
        ("视觉组", "图像 / AI / 感知", COLORS["light_blue"]),
        ("新能源组", "能源系统 / 储能 / 应用", COLORS["deep_blue"]),
        ("运营组", "传播 / 外联 / 开源", COLORS["accent_blue"]),
    ]
    y = 1.35
    for name, desc, color in orgs:
        add_rect(slide, 0.55, y, 0.18, 0.45, color)
        add_textbox(slide, 0.85, y, 1.4, 0.45, name, font_size=12.0, bold=True,
                    color=COLORS["deep_blue"], valign=MSO_ANCHOR.MIDDLE)
        add_textbox(slide, 2.2, y, 4.6, 0.45, desc, font_size=10.0,
                    color=COLORS["gray"], valign=MSO_ANCHOR.MIDDLE)
        y += 0.62
    add_so_what(slide, "团队结构完整，可独立承接从设计、加工到软件与新能源集成的完整项目")
    add_footer(slide, "03")


def build_slide_04(prs, photos):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, SLIDE_WIDTH_IN, SLIDE_HEIGHT_IN, COLORS["ivory"])
    add_page_header(slide, "03", "技术与研发能力", "真实设备、真实工艺、真实调试")
    # Photo
    slide.shapes.add_picture(str(photos["lab_work"]), Inches(0.45), Inches(1.35), width=Inches(6.0), height=Inches(3.6))
    # Capability list
    add_textbox(slide, 7.0, 1.35, 5.7, 0.5, "核心能力", font_size=13.0, bold=True,
                color=COLORS["deep_blue"], valign=MSO_ANCHOR.MIDDLE)
    caps = [
        "3D 打印、激光切割与 CNC 零件快速打样",
        "铝型材框架、碳纤维板与传动机构设计",
        "STM32 / ESP32 电控开发与电机闭环控制",
        "OpenCV 视觉识别与轻量化 AI 部署",
        "新能源电池管理与储能系统验证",
        "开源工具链与跨平台软件交付"
    ]
    add_bullet_list(slide, 7.0, 1.85, 5.7, 3.2, caps, font_size=10.5)
    add_so_what(slide, "从原型到赛场，NEC 具备把概念快速转化为可运行硬件的工程闭环")
    add_footer(slide, "04")


def build_slide_05(prs, photos):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, SLIDE_WIDTH_IN, SLIDE_HEIGHT_IN, COLORS["ivory"])
    add_page_header(slide, "04", "竞赛与荣誉", "国家级赛事持续突破")
    # Big awards
    card1 = add_rounded_rect(slide, 0.55, 1.45, 5.9, 2.0, COLORS["deep_blue"])
    add_textbox(slide, 0.75, 1.6, 5.5, 0.7, "全国大学生机器人大赛", font_size=13.0, bold=True,
                color=COLORS["white"], valign=MSO_ANCHOR.MIDDLE)
    add_textbox(slide, 0.75, 2.3, 5.5, 0.9, "国家级三等奖", font_size=26.0, bold=True,
                color=COLORS["highlight"], valign=MSO_ANCHOR.MIDDLE)
    add_textbox(slide, 0.75, 3.05, 5.5, 0.4, "在 150 余项作品中脱颖而出", font_size=10.0,
                color=COLORS["white"], valign=MSO_ANCHOR.MIDDLE)

    card2 = add_rounded_rect(slide, 6.8, 1.45, 5.9, 2.0, COLORS["accent_blue"])
    add_textbox(slide, 7.0, 1.6, 5.5, 0.7, "AIC 全球人工智能挑战赛", font_size=13.0, bold=True,
                color=COLORS["white"], valign=MSO_ANCHOR.MIDDLE)
    add_textbox(slide, 7.0, 2.3, 5.5, 0.9, "总决赛二等奖", font_size=26.0, bold=True,
                color=COLORS["highlight"], valign=MSO_ANCHOR.MIDDLE)
    add_textbox(slide, 7.0, 3.05, 5.5, 0.4, "人工智能应用与工程落地能力获认可", font_size=10.0,
                color=COLORS["white"], valign=MSO_ANCHOR.MIDDLE)

    # Photo below
    slide.shapes.add_picture(str(photos["robot_team"]), Inches(0.55), Inches(3.75), width=Inches(5.9), height=Inches(2.5))
    slide.shapes.add_picture(str(photos["trae_award"]), Inches(6.8), Inches(3.75), width=Inches(5.9), height=Inches(2.5))

    add_so_what(slide, "竞赛成绩是工程能力的第三方背书，也是政府园区与投资机构评估团队潜力的关键信号")
    add_footer(slide, "05")


def build_slide_06(prs, photos):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, SLIDE_WIDTH_IN, SLIDE_HEIGHT_IN, COLORS["ivory"])
    add_page_header(slide, "05", "社区文化与成长", "从深夜调试到持续迭代")
    # Photo
    slide.shapes.add_picture(str(photos["night_debug"]), Inches(7.2), Inches(1.35), width=Inches(5.7), height=Inches(3.2))
    # Culture bullets
    add_textbox(slide, 0.55, 1.35, 6.2, 0.5, "成长机制", font_size=13.0, bold=True,
                color=COLORS["deep_blue"], valign=MSO_ANCHOR.MIDDLE)
    add_bullet_list(slide, 0.55, 1.85, 6.2, 2.8, [
        "以老带新：高年级成员带领新生完成真实子项目",
        "项目制学习：每赛季围绕赛事目标拆解任务",
        "开源共享：代码、图纸、经验沉淀为社区资产",
        "跨校交流：与产业伙伴、兄弟战队定期联动",
        "失败复盘：每台机器、每场比赛都形成改进清单"
    ], font_size=10.5)
    add_so_what(slide, "社区文化决定了可持续产出能力，赞助商投入的是一支具备自驱迭代能力的年轻团队")
    add_footer(slide, "06")


def build_slide_07(prs, photos):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, SLIDE_WIDTH_IN, SLIDE_HEIGHT_IN, COLORS["ivory"])
    add_page_header(slide, "06", "合作价值主张", "对政府园区与投资机构的意义")
    # Value columns
    values = [
        ("人才池", "提前接触并培养新能源、机器人与嵌入式方向的可用工科人才"),
        ("创新入口", "以高校团队为触角，获取前沿技术原型与低成本验证"),
        ("品牌联动", "与年轻工程师社区绑定，提升企业在高校与开发者群体中的影响力"),
        ("政策衔接", "契合产教融合、双创教育与新能源产业扶持方向"),
    ]
    x = 0.55
    for title, desc in values:
        card = add_rounded_rect(slide, x, 1.45, 2.95, 3.8, COLORS["white"], line_color=COLORS["ivory_dark"], line_width=1)
        add_textbox(slide, x + 0.15, 1.65, 2.65, 0.45, title, font_size=14.0, bold=True,
                    color=COLORS["deep_blue"], align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        add_line(slide, x + 0.35, 2.15, x + 2.6, 2.15, COLORS["highlight"], 1.5)
        add_textbox(slide, x + 0.15, 2.35, 2.65, 2.6, desc, font_size=10.5,
                    color=COLORS["black"], align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, line_spacing=1.3)
        x += 3.15
    add_so_what(slide, "支持 NEC 不仅是资助一支战队，更是以最小成本布局高校新能源创新生态")
    add_footer(slide, "07")


def build_slide_08(prs, photos):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, SLIDE_WIDTH_IN, SLIDE_HEIGHT_IN, COLORS["ivory"])
    add_page_header(slide, "07", "合作方案", "资金 + 场地 + 资源三维支持")
    # Table
    data = [
        ["支持维度", "具体需求", "预期用途", "合作回报"],
        ["资金赞助", "年度研发与赛事经费", "设备、材料、差旅与参赛费用", "冠名权 / 联合品牌露出"],
        ["场地服务", "固定工作室 / 实验室", "加工、调试、会议与日常运营", "空间品牌共建与人才导流"],
        ["产业资源", "导师 / 供应链 / 项目", "技术咨询、样件打样与产业对接", "优先实习与项目转化机会"],
    ]
    add_table(slide, 0.55, 1.45, 4, 4, [1.5, 2.7, 4.0, 3.8], 0.7, data)
    # Process
    add_textbox(slide, 0.55, 4.55, 12.0, 0.45, "合作推进节奏", font_size=13.0, bold=True,
                color=COLORS["deep_blue"], valign=MSO_ANCHOR.MIDDLE)
    steps = [
        ("意向沟通", "1-2 周"),
        ("需求对齐", "2-4 周"),
        ("协议签署", "1-2 周"),
        ("资源注入", "持续"),
        ("季度复盘", "持续"),
    ]
    x = 0.55
    for i, (name, dur) in enumerate(steps):
        add_rounded_rect(slide, x, 5.05, 2.2, 0.65, COLORS["accent_blue"])
        add_textbox(slide, x, 5.05, 2.2, 0.4, name, font_size=11.0, bold=True,
                    color=COLORS["white"], align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        add_textbox(slide, x, 5.4, 2.2, 0.3, dur, font_size=9.0,
                    color=COLORS["white"], align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        if i < len(steps) - 1:
            add_line(slide, x + 2.25, 5.38, x + 2.65, 5.38, COLORS["highlight"], 1.5)
        x += 2.45
    add_so_what(slide, "清晰的资源需求与回报机制，让合作双方的可交付成果从一纸协议起就透明可衡量")
    add_footer(slide, "08")


def build_slide_09(prs, photos):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, SLIDE_WIDTH_IN, SLIDE_HEIGHT_IN, COLORS["ivory"])
    add_page_header(slide, "08", "年度规划与预算", "按季度拆解目标与资金用途")
    # Timeline
    quarters = [
        ("Q1", "招新培训 / 技术栈统一\n新赛季规则研读"),
        ("Q2", "方案设计 / 原型机开发\n关键模块验证"),
        ("Q3", "整机集成 / 联调优化\n区域赛与热身赛"),
        ("Q4", "国赛冲刺 / 成果转化\n年度复盘与招商"),
    ]
    x = 0.55
    for q, desc in quarters:
        add_rounded_rect(slide, x, 1.45, 2.95, 2.4, COLORS["white"], line_color=COLORS["ivory_dark"], line_width=1)
        add_textbox(slide, x, 1.5, 2.95, 0.5, q, font_size=18.0, bold=True,
                    color=COLORS["deep_blue"], align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        add_textbox(slide, x + 0.15, 2.1, 2.65, 1.6, desc, font_size=10.0,
                    color=COLORS["black"], align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.TOP, line_spacing=1.25)
        x += 3.15
    # Budget table (placeholder numbers editable)
    budget_data = [
        ["预算科目", "说明", "备注"],
        ["设备与材料", "电机、传感器、板材、3D 打印耗材", "按实际采购清单填写"],
        ["加工与测试", "激光切割、CNC、场地测试", "按实际发生填写"],
        ["赛事差旅", "全国赛、区域赛交通住宿", "按赛事安排填写"],
        ["运营传播", "宣传物料、开源文档与社区活动", "按实际计划填写"],
    ]
    add_table(slide, 0.55, 4.2, 5, 3, [2.2, 7.0, 2.7], 0.5, budget_data)
    add_textbox(slide, 0.55, 6.8, 12.0, 0.3,
                "注：具体金额需结合赞助规模与年度赛事计划共同确定，本页为可编辑预算框架。",
                font_size=8.0, color=COLORS["gray"], valign=MSO_ANCHOR.MIDDLE)
    add_so_what(slide, " sponsors 的每一笔投入都对应可验证的季度里程碑与公开透明的预算科目")
    add_footer(slide, "09")


def build_slide_10(prs, photos):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, SLIDE_WIDTH_IN, SLIDE_HEIGHT_IN, COLORS["ivory"])
    add_page_header(slide, "09", "生态与资源", "已经建立的外部连接")
    # Photo
    slide.shapes.add_picture(str(photos["opc_event"]), Inches(7.2), Inches(1.35), width=Inches(5.7), height=Inches(3.2))
    # Partners
    add_textbox(slide, 0.55, 1.35, 6.2, 0.5, "已有连接", font_size=13.0, bold=True,
                color=COLORS["deep_blue"], valign=MSO_ANCHOR.MIDDLE)
    partners = [
        "中国电信 OPC 生态合作伙伴交流分享",
        "TRAE Demo Wall 城市人气作品奖（SolarGlyph 等项目）",
        "校内外开源社区与技术社群联动",
        "智能制造、新能源相关企业提供导师与样件支持"
    ]
    add_bullet_list(slide, 0.55, 1.85, 6.2, 2.4, partners, font_size=10.5)
    add_textbox(slide, 0.55, 4.55, 12.0, 0.5, "我们希望引入的伙伴", font_size=13.0, bold=True,
                color=COLORS["deep_blue"], valign=MSO_ANCHOR.MIDDLE)
    targets = [
        "新能源整车 / 动力电池 / 储能企业",
        "机器人与自动化产业链企业",
        "关注硬科技与产教融合的投资机构",
        "地方政府园区与孵化器"
    ]
    add_bullet_list(slide, 0.55, 5.05, 12.0, 1.2, targets, font_size=10.5)
    add_so_what(slide, "NEC 已具备初步产业接口，新赞助商的加入将进一步放大技术转化与人才输送效率")
    add_footer(slide, "10")


def build_slide_11(prs, photos):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, SLIDE_WIDTH_IN, SLIDE_HEIGHT_IN, COLORS["ivory"])
    add_page_header(slide, "10", "我们的工作室", "每天都在发生的工程现场")
    # Photo grid
    slide.shapes.add_picture(str(photos["lab_work"]), Inches(0.55), Inches(1.35), width=Inches(3.95), height=Inches(2.22))
    slide.shapes.add_picture(str(photos["robot_team"]), Inches(4.75), Inches(1.35), width=Inches(3.95), height=Inches(2.22))
    slide.shapes.add_picture(str(photos["night_debug"]), Inches(8.95), Inches(1.35), width=Inches(3.95), height=Inches(2.22))
    # Labels
    add_textbox(slide, 0.55, 3.65, 3.95, 0.35, "日常研发", font_size=11.0, bold=True,
                color=COLORS["deep_blue"], align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_textbox(slide, 4.75, 3.65, 3.95, 0.35, "机器调试", font_size=11.0, bold=True,
                color=COLORS["deep_blue"], align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_textbox(slide, 8.95, 3.65, 3.95, 0.35, "深夜协作", font_size=11.0, bold=True,
                color=COLORS["deep_blue"], align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    # Text
    add_textbox(slide, 0.55, 4.25, 12.0, 1.2,
                "这些照片记录了 NEC 成员真实的研发日常：从图纸到零件、从代码到整机、从白天到深夜。\n"
                "我们需要的不是更好的拍摄环境，而是能够长期稳定运行的场地与资源。",
                font_size=11.0, color=COLORS["black"], valign=MSO_ANCHOR.TOP, line_spacing=1.3)
    add_so_what(slide, "真实的工程现场胜过千言万语，场地与资金支持将直接决定下一台机器能否按时站上赛场")
    add_footer(slide, "11")


def build_slide_12(prs, photos):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, SLIDE_WIDTH_IN, SLIDE_HEIGHT_IN, COLORS["deep_blue"])
    add_textbox(slide, 0.55, 2.2, 12.0, 1.0, "期待与您共建", font_size=38.0, bold=True,
                color=COLORS["white"], align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_textbox(slide, 0.55, 3.3, 12.0, 0.6, "新能源开发者社区 × 政府园区 × 投资机构", font_size=16.0,
                color=COLORS["highlight"], align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_textbox(slide, 0.55, 4.2, 12.0, 1.0,
                "联系人：[待填写]\n邮箱：[待填写]\n地址：常州工学院\n社区：NEC 新能源开发者社区",
                font_size=12.0, color=COLORS["white"], align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.TOP, line_spacing=1.4)
    add_textbox(slide, 0.55, 7.05, 12.0, 0.3, "常州工学院 NEC 新能源开发者社区 · 企业合作提案 · 2026",
                font_size=7.0, color=COLORS["white"], align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    photos = ensure_assets()

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_WIDTH_IN)
    prs.slide_height = Inches(SLIDE_HEIGHT_IN)

    build_slide_01(prs, photos)
    build_slide_02(prs, photos)
    build_slide_03(prs, photos)
    build_slide_04(prs, photos)
    build_slide_05(prs, photos)
    build_slide_06(prs, photos)
    build_slide_07(prs, photos)
    build_slide_08(prs, photos)
    build_slide_09(prs, photos)
    build_slide_10(prs, photos)
    build_slide_11(prs, photos)
    build_slide_12(prs, photos)

    prs.save(PPTX_PATH)
    print(f"Saved PPTX: {PPTX_PATH}")


if __name__ == "__main__":
    main()
